"""把配方 / 材料导出为 Excel（仅文本，不嵌入图片）。"""

from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session, selectinload

from app.core.database import Ingredient, Recipe, RecipeIngredient

HEADERS = [
    "来源",
    "源站ID",
    "中文名",
    "英文名",
    "分类",
    "基酒",
    "杯型",
    "技法",
    "技法展示",
    "酒精度",
    "冰型",
    "装饰",
    "风味标签",
    "简介",
    "英文简介",
    "材料",
    "步骤",
    "图片URL",
]
INGREDIENT_HEADERS = ["材料名", "英文名", "分类", "主单位", "出现次数", "图片URL"]


def _format_amount(amount) -> str:
    if amount is None:
        return ""
    return f"{float(amount):g}"


def _ingredient_text(recipe: Recipe) -> str:
    parts = []
    for item in recipe.ingredients:
        master = item.ingredient
        name = (master.name if master else "") or ""
        parts.append(f"{name} {_format_amount(item.amount)}{item.unit or ''}".strip())
    return "；".join(parts)


def build_recipes_workbook(db: Session) -> BytesIO:
    recipes = (
        db.query(Recipe)
        .options(selectinload(Recipe.ingredients).selectinload(RecipeIngredient.ingredient))
        .order_by(Recipe.source, Recipe.id)
        .all()
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "配方"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.row_dimensions[1].height = 22
    wrap = Alignment(wrap_text=True, vertical="top")

    for index, recipe in enumerate(recipes, start=2):
        try:
            tags = json.loads(recipe.flavor_tags or "[]")
        except json.JSONDecodeError:
            tags = []
        sheet.append(
            [
                recipe.source,
                recipe.source_id,
                recipe.name,
                recipe.name_en,
                recipe.category,
                recipe.base_spirit,
                recipe.glass,
                recipe.method,
                recipe.method_label,
                recipe.abv_range,
                recipe.ice_type,
                recipe.garnish,
                "、".join(tags) if isinstance(tags, list) else "",
                recipe.description,
                recipe.description_en,
                _ingredient_text(recipe),
                recipe.instructions,
                recipe.image_url or "",
            ]
        )
        for cell in sheet[index]:
            cell.alignment = wrap

    widths = {
        "A": 12,
        "C": 20,
        "D": 22,
        "N": 36,
        "P": 40,
        "Q": 40,
        "R": 48,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for col_idx in range(1, len(HEADERS) + 1):
        letter = get_column_letter(col_idx)
        if letter not in widths:
            sheet.column_dimensions[letter].width = 14

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_ingredients_workbook(db: Session) -> BytesIO:
    """导出材料主表，附带在配方中的出现次数。"""
    usage_counts = dict(
        db.query(RecipeIngredient.ingredient_id, func.count(RecipeIngredient.id))
        .group_by(RecipeIngredient.ingredient_id)
        .all()
    )
    rows = db.query(Ingredient).order_by(Ingredient.name, Ingredient.id).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "材料清单"
    sheet.append(INGREDIENT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 10
    sheet.column_dimensions["E"].width = 12
    sheet.column_dimensions["F"].width = 56

    items = sorted(
        [
            {
                "name": row.name or "",
                "name_en": row.name_en or "",
                "category": row.category or "",
                "primary_unit": row.primary_unit or "",
                "count": int(usage_counts.get(row.id) or 0),
                "image_url": row.image_url or "",
            }
            for row in rows
        ],
        key=lambda item: (-int(item["count"]), str(item["name"])),
    )
    for item in items:
        sheet.append(
            [
                item["name"],
                item["name_en"],
                item["category"],
                item["primary_unit"],
                item["count"],
                item["image_url"],
            ]
        )

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
